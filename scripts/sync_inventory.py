#!/usr/bin/env python3
"""Build data/inventory.csv from the Dropbox Excel sheet (only the public columns).

Reads: name, genus/type, sell price, in stock  (header text is matched loosely, see ALIASES).
Descriptions: extras.csv text first, else the genus text in data/genus_descriptions.csv, else (if ANTHROPIC_API_KEY is set) a one-sentence AI description cached in data/descriptions_cache.json.
Merges data/extras.csv (matched by plant name) for things the sheet doesn't have:
sku, genus (only used when the sheet's Genus/Type is blank), category, size, description, image (photos separated by |), status, sale_price.
Anything else in the spreadsheet (costs, notes, customers...) is never read or published.

Usage:
  DROPBOX_INVENTORY_URL=<share link> python3 scripts/sync_inventory.py
  python3 scripts/sync_inventory.py --file local.xlsx        # test with a local copy
  add --dry to print a summary without writing.
Optional env: SHEET_NAME (tab to read; default: first tab with the right headers).
"""
import json, csv, io, os, re, sys, urllib.request
from pathlib import Path
from urllib.parse import urlparse, urlunparse, parse_qsl, urlencode
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT, EXTRAS, PHOTOS = ROOT / "data/inventory.csv", ROOT / "data/extras.csv", ROOT / "data/photos.csv"
GENUS_DESC, DESC_CACHE = ROOT / "data/genus_descriptions.csv", ROOT / "data/descriptions_cache.json"
ALIASES = {
    "name": ["name", "plant", "plant name", "item"],
    "genus": ["genus/type", "genus / type", "genus", "type"],
    "price": ["sell price", "selling price", "price", "sale price", "retail"],
    "stock": ["in stock", "stock", "qty", "quantity", "count"],
}
COLS = ["sku", "name", "genus", "category", "price", "sale_price", "status", "stock", "size", "description", "image"]
MAX_ROWS = 3000

norm = lambda s: re.sub(r"[^a-z0-9/ ]+", "", str(s or "").strip().lower())
key = lambda s: re.sub(r"\s+", " ", str(s or "").strip().lower())
def slug(s): return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")[:60]
def clean(s): return re.sub(r"[\x00-\x1f\x7f]", " ", str(s if s is not None else "")).strip()

def fetch(url):
    u = urlparse(url)
    if "dropbox.com" not in u.netloc: sys.exit("DROPBOX_INVENTORY_URL must be a dropbox.com share link")
    q = dict(parse_qsl(u.query)); q["dl"] = "1"
    req = urllib.request.Request(urlunparse(u._replace(query=urlencode(q))), headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r: data = r.read(30_000_000)
    if data[:2] != b"PK": sys.exit("Downloaded file is not an .xlsx (check the share link permissions)")
    return io.BytesIO(data)

def find_table(wb, want_sheet):
    for ws in wb.worksheets:
        if want_sheet and ws.title.strip().lower() != want_sheet.strip().lower(): continue
        rows = list(ws.iter_rows(min_row=1, max_row=MAX_ROWS, values_only=True))
        for i, row in enumerate(rows[:25]):
            heads = [norm(c) for c in row]
            idx = {}
            for field, names in ALIASES.items():
                for n in names:
                    if n in heads: idx[field] = heads.index(n); break
            if "name" in idx and ("price" in idx or "stock" in idx):
                return ws.title, idx, rows[i + 1:]
    sys.exit(f"Could not find a header row with name / price / stock columns{' in sheet ' + repr(want_sheet) if want_sheet else ''}")

def num(v):
    if v is None or str(v).strip() == "": return None
    m = re.search(r"-?\d+(?:\.\d+)?", str(v).replace(",", ""))
    return float(m.group()) if m else None

def stock_val(v):
    s = str(v if v is not None else "").strip().lower()
    if s == "": return ""
    if s in ("no", "n", "out", "sold", "sold out", "none"): return "0"
    n = num(s)
    return "" if n is None else str(max(0, int(n)))

def load_photos():
    """{normalized filename: image path} from data/photos.csv (written by sync_photos.py)."""
    out = {}
    if PHOTOS.exists():
        with open(PHOTOS, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f): out[r["key"]] = r["image"]
    return out

def _nk(s): return re.sub(r"[^a-z0-9]", "", str(s or "").lower())

SIZES = ("large", "small", "medium")
def canon(key):
    """Reduce a name to (core, size, variegated?) so abbreviations, word order and optional size words don't matter:
    'vrg'/'var'/'variegata' -> variegated, 'lrg' -> large, 'sml' -> small; 'Semi Glabrata Large' == 'large semi glabrata'."""
    k = _nk(key)
    k = re.sub(r"vrgt?|varieg(?:ated|ata)?|var$", "variegated", k)
    k = re.sub(r"lrg|lge", "large", k); k = re.sub(r"sml", "small", k); k = re.sub(r"med$", "medium", k)
    size = next((z for z in SIZES if z in k), None)
    for z in SIZES: k = k.replace(z, "")
    var = "variegated" in k
    return k.replace("variegated", ""), size, var

ABBR = {"vrg": "variegated", "vrgt": "variegated", "var": "variegated", "vari": "variegated", "varieg": "variegated", "variegata": "variegated",
        "lrg": "large", "lge": "large", "sml": "small", "med": "medium"}
def tokens(label):
    return [ABBR.get(t, t) for t in re.findall(r"[a-z]+|\d+", str(label or "").lower())]

def load_aliases():
    """data/photo_aliases.csv: 'photo' (file name) -> 'plant' (exact name in the sheet), for photos whose name can't be matched automatically."""
    f = ROOT / "data/photo_aliases.csv"; out = {}
    if f.exists():
        with open(f, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                if r.get("photo") and r.get("plant"): out[_nk(r["photo"])] = _nk(r["plant"])
    return out

def load_labels():
    f = ROOT / "data/photos_labels.json"
    try: return json.loads(f.read_text(encoding="utf-8")) if f.exists() else {}
    except Exception: return {}

def assign_photos(photos, plants, labels=None):
    """Match photo files to plants. Ignores capitals/spaces/punctuation, allows a trailing number for extra
    photos ('Name 2'), an optional genus in front ('Adenium Multiflorum'), abbreviations (vrg = variegated),
    word order and size words, and small misspellings. plants = [(name, genus)]; returns {plant index: [(order, path, filekey)]}."""
    from difflib import SequenceMatcher
    info = []
    for name, genus in plants:
        core, size, var = canon(name); g = _nk(genus)
        forms = {core, g + core} if g else {core}
        if g:   # the sheet's "genus/type" column is sometimes part of the name (e.g. Jade Dwarf + Variegated striata)
            c2, s2, v2 = canon(name + " " + genus); forms.add(c2); size = size or s2; var = var or v2
        info.append((forms, size, var, {frozenset(tokens(name)), frozenset(tokens(name) + tokens(genus))}))
    out = {}
    aliases = load_aliases()
    for fk, path in photos.items():
        al = aliases.get(fk) or (aliases.get(re.sub(r"\d{1,2}$", "", fk)) if re.search(r"\d{1,2}$", fk) else None)
        if al:   # you told us which plant this photo is for
            order = int(re.search(r"(\d{1,2})$", fk).group(1)) if (not aliases.get(fk) and re.search(r"\d{1,2}$", fk)) else 1
            hit = [pi for pi, (n, _g) in enumerate(plants) if _nk(n) == al]
            if hit:
                for pi in hit: out.setdefault(pi, []).append((order, path, fk))
                continue
        bases = [(fk, 1)]
        m = re.fullmatch(r"(.+?)(\d{1,2})", fk)
        if m: bases.append((m.group(1), int(m.group(2))))
        best = None   # (score, [plant indexes], order)
        for base, order in bases:
            core, size, var = canon(base)
            scores = []
            for pi, (forms, psize, pvar, ptok) in enumerate(info):
                if size and psize and size != psize: continue                  # 'large' vs 'small': never the same plant
                if var and not pvar: continue                                  # a variegated photo never goes to a plain plant
                tk = [t for t in tokens((labels or {}).get(fk)) ] if labels and fk in labels else []
                if order > 1 and tk and tk[-1].isdigit(): tk = tk[:-1]
                if tk and frozenset(tk) in ptok:      # same words in any order (Jade Dwarf Vrg Striata == Jade striata Dwarf Variegated)
                    scores.append((1.0, pi)); continue
                sc = 1.0 if core in forms else (max(SequenceMatcher(None, core, f).ratio() for f in forms) if len(core) >= 5 else 0)
                if sc < 0.84: continue
                if size != psize: sc *= 0.97                                   # size given on only one side
                if var != pvar: sc *= 0.93                                     # plant is variegated, photo name doesn't say so
                scores.append((sc, pi))
            if not scores: continue
            scores.sort(reverse=True)
            top = scores[0][0]
            ties = [p for sc, p in scores if abs(sc - top) < 1e-9]
            if top < 0.995 and len(ties) == 1 and len(scores) > 1 and scores[1][0] > top - 0.04: continue   # ambiguous guess
            if best is None or top > best[0]: best = (top, ties, order)
        if best:
            for pi in best[1]: out.setdefault(pi, []).append((best[2], path, fk))
    for v in out.values(): v.sort()
    return out

def load_genus_desc():
    out = {}
    if GENUS_DESC.exists():
        with open(GENUS_DESC, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f): out[key(r["genus"])] = r["description"].strip()
    return out

def ai_description(name, genus, cache):
    """Optional: one-sentence description from the Claude API (needs ANTHROPIC_API_KEY). Cached so each plant is only written once."""
    ck = key(name) + "|" + key(genus)
    if ck in cache: return cache[ck]
    api = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if not api: return ""
    import json
    body = json.dumps({"model": os.environ.get("DESC_MODEL", "claude-haiku-5-5"), "max_tokens": 120, "messages": [{"role": "user", "content":
        f"Write ONE short factual sentence (max 25 words) describing this cactus/succulent for a plant shop: '{name}'" + (f" (genus {genus})" if genus else "") +
        ". Describe only the plant itself, no care advice, no price, no marketing. If you are unsure of the specific plant, describe its genus or general type. Output only the sentence."}]}).encode()
    req = urllib.request.Request("https://api.anthropic.com/v1/messages", data=body, headers={"x-api-key": api, "anthropic-version": "2023-06-01", "content-type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=40) as r: text = json.loads(r.read())["content"][0]["text"].strip()
    except Exception as e:
        print("description API skipped for", name, type(e).__name__); return ""
    text = clean(text)[:240]; cache[ck] = text; return text

def main():
    args = sys.argv[1:]
    dry = "--dry" in args
    if "--file" in args: src = open(args[args.index("--file") + 1], "rb")
    else:
        url = os.environ.get("DROPBOX_INVENTORY_URL", "").strip()
        if not url: sys.exit("Set DROPBOX_INVENTORY_URL (or use --file)")
        src = fetch(url)
    wb = load_workbook(src, read_only=True, data_only=True)
    title, idx, rows = find_table(wb, os.environ.get("SHEET_NAME", ""))
    extras = {}
    if EXTRAS.exists():
        with open(EXTRAS, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f): extras[key(r.get("name"))] = r
    photos, claimed = load_photos(), set()
    plist = []
    for row in rows:
        if not row or idx["name"] >= len(row): continue
        nm = clean(row[idx["name"]])
        if not nm or re.fullmatch(r"(grand )?totals?:?", key(nm)): continue
        if "stock" in idx and idx["stock"] < len(row) and stock_val(row[idx["stock"]]) == "0": continue   # sold out / 0 = removed from the site
        gv = clean(row[idx["genus"]]) if "genus" in idx and idx["genus"] < len(row) else ""
        plist.append((nm, gv or extras.get(key(nm), {}).get("genus", "")))
    matched = assign_photos(photos, plist, load_labels())
    pn = -1
    gdesc = load_genus_desc()
    import json
    cache = json.loads(DESC_CACHE.read_text()) if DESC_CACHE.exists() else {}
    ai_n = 0
    out, used = [], set()
    for row in rows:
        if not row or idx["name"] >= len(row): continue
        name = clean(row[idx["name"]])
        if not name or re.fullmatch(r"(grand )?totals?:?", key(name)): continue   # skip blanks and the sheet's Total row
        if "stock" in idx and idx["stock"] < len(row) and stock_val(row[idx["stock"]]) == "0": continue   # 0 in stock: plant is deleted from the site
        pn += 1
        g = lambda f: row[idx[f]] if f in idx and idx[f] < len(row) else None
        price = num(g("price"))
        ex = extras.get(key(name), {})
        genus = clean(g("genus")) or ex.get("genus", "")   # sheet wins; extras.csv fills blanks
        sku = ex.get("sku") if (ex.get("sku") and ex.get("sku") not in used) else slug(name)
        n = 2
        while sku in used: sku = f"{slug(name)}-{n}"; n += 1
        used.add(sku)
        out.append({"sku": sku, "name": name, "genus": genus, "category": ex.get("category") or genus or "Plants",
                    "price": ("%g" % price) if price else "", "sale_price": ex.get("sale_price", ""),
                    "status": ex.get("status") or "In Stock", "stock": stock_val(g("stock")),
                    "size": ex.get("size", ""), "description": ex.get("description") or gdesc.get(key(genus), "") or ai_description(name, genus, cache), "image": "|".join(p for _, p, _ in matched.get(pn, [])) or ex.get("image", "")})
        for _, _, fk in matched.get(pn, []): claimed.add(fk)
        if len(out) >= MAX_ROWS: break
    if not out: sys.exit("No plants found; leaving inventory.csv unchanged")
    print(f"sheet '{title}': {len(out)} plants, total stock {sum(int(r['stock'] or 0) for r in out)} ({sum(1 for r in out if r['image'])} with photos, {sum(1 for r in out if not r['price'])} without a price)")
    loose = sorted(set(photos) - claimed)
    if loose: print("photos that match no plant name (removed from the site):", ", ".join(loose[:15]))
    if loose and not dry and len(claimed) > 0:
        # A photo whose plant is gone from the sheet (or has 0 in stock) is deleted from the site. (Dropbox is untouched.)
        ph_dir = ROOT / "assets/photos"
        for k in loose:
            for p in (ph_dir / f"photo-{k}.jpg", ph_dir / "web" / f"photo-{k}.webp", ph_dir / "thumb" / f"photo-{k}.webp"):
                if p.exists(): p.unlink()
        keep = [(k, v) for k, v in photos.items() if k not in loose]
        with open(PHOTOS, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n"); w.writerow(["key", "image"])
            for k, v in sorted(keep): w.writerow([k, v])
    if cache and not dry: DESC_CACHE.write_text(json.dumps(cache, indent=0, sort_keys=True, ensure_ascii=False))
    if dry: return
    buf = io.StringIO(); w = csv.DictWriter(buf, fieldnames=COLS, lineterminator="\n"); w.writeheader(); w.writerows(out)
    new = buf.getvalue()
    if OUT.exists() and OUT.read_text(encoding="utf-8") == new: print("no changes"); return
    OUT.write_text(new, encoding="utf-8"); print("wrote", OUT.relative_to(ROOT))

if __name__ == "__main__": main()
